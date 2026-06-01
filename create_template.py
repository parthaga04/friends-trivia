"""Run this script once to generate friends_trivia_questions.xlsx.

Fill in the template with your own questions before game night.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


PURPLE = 'FF6B2D8B'
GOLD   = 'FFFFD700'
DARK   = 'FF1A0B36'
WHITE  = 'FFFFFFFF'
LIGHT_PURPLE = 'FFD7BDE2'


def header_style(cell, bg=PURPLE, fg=WHITE):
    cell.font = Font(bold=True, color=fg, size=12)
    cell.fill = PatternFill('solid', fgColor=bg)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin = Side(style='thin', color='FFAAAAAA')
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def data_style(cell, row_num):
    bg = 'FFF5EEF8' if row_num % 2 == 0 else 'FFFFFFFF'
    cell.fill = PatternFill('solid', fgColor=bg)
    cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    thin = Side(style='thin', color='FFDDDDDD')
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def set_col_widths(ws, widths):
    for col, w in zip('ABCDEFGH', widths):
        ws.column_dimensions[col].width = w


def create_template():
    wb = openpyxl.Workbook()

    # ── Instructions sheet ────────────────────────────────────────────────
    ws_info = wb.active
    ws_info.title = 'Instructions'
    ws_info.column_dimensions['A'].width = 80

    lines = [
        ('FRIENDS TRIVIA – Question Template', True, 18, PURPLE),
        ('', False, 11, DARK),
        ('HOW TO FILL IN THIS FILE', True, 13, PURPLE),
        ('', False, 11, DARK),
        ('1.  "Main Board" sheet  –  fill in your 40 questions (4 categories × 10 questions each).', False, 11, DARK),
        ('    • Category  : any name you choose (e.g. "Most Likely To", "Hot Takes", "Childhood", "Secrets")', False, 11, DARK),
        ('    • Person    : the EXACT name of the player this question is about.', False, 11, DARK),
        ('                  It must match the name you enter for that player at game setup.', False, 11, DARK),
        ('    • Question  : write the question as it will be read aloud.', False, 11, DARK),
        ('    • Answer    : the correct answer that will be revealed on screen.', False, 11, DARK),
        ('', False, 11, DARK),
        ('2.  "Lightning Round" sheet  –  fill in rapid-fire questions (aim for 15–25 per team pair).', False, 11, DARK),
        ('    • Same format: Person | Question | Answer', False, 11, DARK),
        ('    • Shorter questions and answers work best here.', False, 11, DARK),
        ('', False, 11, DARK),
        ('3.  Save the file and upload it on the game setup page.', False, 11, DARK),
        ('', False, 11, DARK),
        ('IMPORTANT', True, 12, 'FFCC0000'),
        ('Player names in the "Person" column must exactly match the names you type at setup.', False, 11, DARK),
        ('The game routes Team B\'s questions to players on Team A, and vice versa.', False, 11, DARK),
    ]

    for i, (text, bold, size, color) in enumerate(lines, start=1):
        cell = ws_info.cell(row=i, column=1, value=text)
        cell.font = Font(bold=bold, size=size, color=color)
        cell.alignment = Alignment(wrap_text=True)

    # ── Main Board sheet ──────────────────────────────────────────────────
    ws_main = wb.create_sheet('Main Board')
    headers = ['Category', 'Person', 'Question', 'Answer']
    for col, h in enumerate(headers, start=1):
        header_style(ws_main.cell(row=1, column=col, value=h))

    ws_main.row_dimensions[1].height = 28
    set_col_widths(ws_main, [20, 18, 55, 35])

    sample_main = [
        # Category,            Person,    Question,                                                        Answer
        ('Most Likely To',   '[Name A]', 'Who is most likely to show up two hours late to their own party?', 'Because they got distracted shopping'),
        ('Most Likely To',   '[Name B]', 'Who is most likely to cry at a commercial?',                       'Every single time'),
        ('Most Likely To',   '[Name A]', 'Who is most likely to order the most expensive thing on the menu when someone else is paying?', 'Without hesitation'),
        ('Most Likely To',   '[Name B]', 'Who is most likely to start a conversation with a complete stranger?', 'They cannot help themselves'),
        ('Hot Takes',        '[Name A]', 'What is [Name A]\'s most controversial food opinion?',             'Replace with their actual opinion'),
        ('Hot Takes',        '[Name B]', 'What does [Name B] secretly think about reality TV shows?',        'Replace with their actual opinion'),
        ('Hot Takes',        '[Name A]', 'If [Name A] could only eat one food for the rest of their life, what would they choose?', 'Replace with the real answer'),
        ('Hot Takes',        '[Name B]', 'What is the one thing [Name B] refuses to do, no matter what?',    'Replace with the real answer'),
        ('Childhood',        '[Name A]', 'What was [Name A]\'s most embarrassing moment in high school?',    'Replace with the real story'),
        ('Childhood',        '[Name B]', 'What nickname did [Name B] have growing up?',                      'Replace with the real nickname'),
        ('Childhood',        '[Name A]', 'What was [Name A]\'s first job?',                                  'Replace with the real answer'),
        ('Childhood',        '[Name B]', 'What sport or hobby did [Name B] try as a kid and quit after one lesson?', 'Replace with the real answer'),
        ('Dark Secrets',     '[Name A]', 'What guilty pleasure show does [Name A] binge-watch alone?',       'Replace with the real answer'),
        ('Dark Secrets',     '[Name B]', 'What talent does [Name B] have that almost no one knows about?',   'Replace with the real answer'),
        ('Dark Secrets',     '[Name A]', 'What is the most embarrassing song on [Name A]\'s playlist?',      'Replace with the real answer'),
        ('Dark Secrets',     '[Name B]', 'What habit does [Name B] have that they think no one notices?',    'Replace with the real answer'),
    ]

    for row_num, row_data in enumerate(sample_main, start=2):
        for col, val in enumerate(row_data, start=1):
            cell = ws_main.cell(row=row_num, column=col, value=val)
            data_style(cell, row_num)
        ws_main.row_dimensions[row_num].height = 40

    ws_main.freeze_panes = 'A2'

    # ── Lightning Round sheet ─────────────────────────────────────────────
    ws_light = wb.create_sheet('Lightning Round')
    light_headers = ['Person', 'Question', 'Answer']
    for col, h in enumerate(light_headers, start=1):
        header_style(ws_light.cell(row=1, column=col, value=h))

    ws_light.row_dimensions[1].height = 28
    set_col_widths(ws_light, [18, 55, 30])

    sample_lightning = [
        ('[Name A]', 'What is [Name A]\'s go-to karaoke song?',           'Replace with real answer'),
        ('[Name B]', 'What does [Name B] always order at a restaurant?',   'Replace with real answer'),
        ('[Name A]', 'What is [Name A]\'s biggest fear?',                  'Replace with real answer'),
        ('[Name B]', 'What was [Name B]\'s first car?',                    'Replace with real answer'),
        ('[Name A]', 'What is [Name A]\'s go-to excuse when cancelling plans?', 'Replace with real answer'),
        ('[Name B]', 'What is [Name B]\'s most-used emoji?',               'Replace with real answer'),
        ('[Name A]', 'What does [Name A] always say when they\'re nervous?', 'Replace with real answer'),
        ('[Name B]', 'How many alarms does [Name B] set in the morning?',  'Replace with real answer'),
    ]

    for row_num, row_data in enumerate(sample_lightning, start=2):
        for col, val in enumerate(row_data, start=1):
            cell = ws_light.cell(row=row_num, column=col, value=val)
            data_style(cell, row_num)
        ws_light.row_dimensions[row_num].height = 36

    ws_light.freeze_panes = 'A2'

    outfile = 'friends_trivia_questions.xlsx'
    wb.save(outfile)
    print(f'Template created: {outfile}')
    print('Open it, fill in your questions, and upload it at game setup!')


if __name__ == '__main__':
    create_template()
