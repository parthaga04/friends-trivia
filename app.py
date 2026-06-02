from flask import Flask, render_template, jsonify, request, redirect, url_for, Response
import openpyxl
import os
import random
import hashlib
import requests as http_requests

app = Flask(__name__)

GAME_DATA = {}

# ── ElevenLabs TTS ────────────────────────────────────────────
ELEVENLABS_API_KEY = '28bd4cf41a1816ce2add4974e0f64a631e1ed594447f09fa4f1d0604df7200e4'
ELEVENLABS_VOICE_ID = 'LwVNjJUMuxV61q2kUIqx'

TTS_CACHE = {}


@app.route('/api/tts', methods=['POST'])
def tts():
    text = (request.json or {}).get('text', '').strip()
    if not text:
        return '', 400

    cache_key = hashlib.md5(text.encode()).hexdigest()
    if cache_key in TTS_CACHE:
        return Response(TTS_CACHE[cache_key], mimetype='audio/mpeg')

    if not ELEVENLABS_API_KEY or not ELEVENLABS_VOICE_ID or ELEVENLABS_VOICE_ID == 'PASTE_VOICE_ID_HERE':
        return '', 503

    r = http_requests.post(
        f'https://api.elevenlabs.io/v1/text-to-speech/{ELEVENLABS_VOICE_ID}',
        headers={'xi-api-key': ELEVENLABS_API_KEY, 'Content-Type': 'application/json'},
        json={
            'text': text,
            'model_id': 'eleven_turbo_v2',
            'voice_settings': {'stability': 0.5, 'similarity_boost': 0.75}
        },
        timeout=10
    )
    if r.status_code != 200:
        return '', 503

    TTS_CACHE[cache_key] = r.content
    return Response(r.content, mimetype='audio/mpeg')


def shuffle(lst):
    lst = list(lst)
    random.shuffle(lst)
    return lst


def parse_excel(filepath):
    wb = openpyxl.load_workbook(filepath)

    questions_by_category = {}

    if 'Main Board' in wb.sheetnames:
        ws = wb['Main Board']
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0] or not row[1] or not row[2]:
                continue
            category = str(row[0]).strip()
            person = str(row[1]).strip()
            question_text = str(row[2]).strip()
            answer = str(row[3]).strip() if row[3] else ''

            q = {'person': person, 'question': question_text, 'answer': answer}

            if category not in questions_by_category:
                questions_by_category[category] = []
            questions_by_category[category].append(q)

    for cat in questions_by_category:
        questions_by_category[cat] = shuffle(questions_by_category[cat])

    lightning_questions = []
    if 'Lightning Round' in wb.sheetnames:
        ws = wb['Lightning Round']
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0] or not row[1]:
                continue
            lightning_questions.append({
                'person': str(row[0]).strip(),
                'question': str(row[1]).strip(),
                'answer': str(row[2]).strip() if row[2] else ''
            })

    return questions_by_category, shuffle(lightning_questions)


@app.route('/')
def index():
    global GAME_DATA
    GAME_DATA = {}
    return render_template('index.html')


@app.route('/start', methods=['POST'])
def start_game():
    global GAME_DATA

    team_a_name = request.form.get('team_a_name', 'Team A').strip() or 'Team A'
    team_b_name = request.form.get('team_b_name', 'Team B').strip() or 'Team B'
    lightning_mode = request.form.get('lightning_mode', 'tie')
    num_questions = int(request.form.get('num_questions', 10))

    players_a = [p.strip() for p in request.form.getlist('players_a') if p.strip()]
    players_b = [p.strip() for p in request.form.getlist('players_b') if p.strip()]

    if 'excel_file' not in request.files:
        return redirect(url_for('index'))
    file = request.files['excel_file']
    if not file.filename:
        return redirect(url_for('index'))

    filepath = '/tmp/friends_trivia_game.xlsx'
    file.save(filepath)

    questions_by_category, lightning_questions = parse_excel(filepath)

    GAME_DATA = {
        'teams': [team_a_name, team_b_name],
        'players': [players_a, players_b],
        'questions': questions_by_category,
        'lightning': lightning_questions,
        'lightning_mode': lightning_mode,
        'num_questions': num_questions,
    }

    return redirect(url_for('game'))


@app.route('/game')
def game():
    if not GAME_DATA:
        return redirect(url_for('index'))
    return render_template('game.html')


@app.route('/api/game-data')
def get_game_data():
    return jsonify(GAME_DATA)


if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=True)
